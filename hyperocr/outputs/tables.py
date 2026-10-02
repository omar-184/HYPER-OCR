"""Each table as its own Word document (.docx), merged cells kept."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..document import mostly_rtl
from .tablegrid import parse_table

# Child order required by the OOXML schema (Word rejects files that break it).
RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
             "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w",
             "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign",
             "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
             "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
             "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
               "tblCaption", "tblDescription", "tblPrChange"]


def _put(parent, name: str, order: list[str], val: str | None = None):
    """Add w:<name> to `parent` at its schema position (replacing one already there)."""
    tag = qn("w:" + name)
    existing = parent.find(tag)
    if existing is not None:
        parent.remove(existing)
    el = OxmlElement("w:" + name)
    if val is not None:
        el.set(qn("w:val"), val)
    rank = order.index(name)
    for child in parent:
        local = child.tag.rsplit("}", 1)[-1]
        if local in order and order.index(local) > rank:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


LABELS = {
    "en": {"table": "Table {n}", "page": "page {p}", "scan": "Original scan", "source": "Source: {file}, page {p}",
           "check": "Read by OCR and not verified. Check every number against the picture of the original table "
                    "below before you use it."},
    "ar": {"table": "جدول {n}", "page": "صفحة {p}", "scan": "الصورة الأصلية", "source": "المصدر: {file}، صفحة {p}",
           "check": "قُرئ هذا الجدول آلياً ولم يُتحقَّق منه. طابِق كل رقم مع صورة الجدول الأصلي أدناه قبل استخدامه."},
}
FONT = "Arial"   # has Latin and Arabic glyphs on Windows, macOS and LibreOffice


def write_table_docx(
    path: Path,
    table_html: str,
    number: int,
    page_no: int,
    source_name: str,
    caption: str = "",
    snapshot_png: bytes | None = None,
    ui_lang: str = "en",
) -> None:
    labels = LABELS.get(ui_lang, LABELS["en"])
    grid = parse_table(table_html)
    rtl = grid.rtl or mostly_rtl(" ".join(c.text for c in grid.cells))
    doc = Document()
    _base_font(doc)
    section = doc.sections[0]
    if grid.cols > 6:   # wide tables read better across the page
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    usable = section.page_width - section.left_margin - section.right_margin

    title_rtl = ui_lang == "ar"
    title = doc.add_heading(level=2)
    _run(title, "%s · %s" % (labels["table"].format(n=number), labels["page"].format(p=page_no)), title_rtl)
    _para_dir(title, title_rtl)
    if caption:
        p = doc.add_paragraph()
        _run(p, caption, mostly_rtl(caption), italic=True)
        _para_dir(p, mostly_rtl(caption))
    p = doc.add_paragraph()
    _run(p, labels["check"], title_rtl, bold=True, size=9.5)
    _para_dir(p, title_rtl)

    if grid.rows and grid.cols:
        table = doc.add_table(rows=grid.rows, cols=grid.cols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        if rtl:
            _put(table._tbl.tblPr, "bidiVisual", TBLPR_ORDER)  # first column on the right
        header_rows = {c.row for c in grid.cells if c.header} or ({0} if grid.rows > 1 else set())
        for cell in grid.cells:
            target = table.cell(cell.row, cell.col)
            last = table.cell(min(grid.rows, cell.row + cell.rowspan) - 1, min(grid.cols, cell.col + cell.colspan) - 1)
            if last is not target:
                target = target.merge(last)
            para = target.paragraphs[0]
            cell_rtl = mostly_rtl(cell.text) if cell.text else rtl
            lines = cell.text.split("\n") if cell.text else [""]
            for i, line in enumerate(lines):
                if i:
                    para = target.add_paragraph()
                _run(para, line, cell_rtl, bold=cell.row in header_rows)
                _para_dir(para, cell_rtl)
        for r in sorted(header_rows):
            if r < grid.rows:
                tr_pr = table.rows[r]._tr.get_or_add_trPr()
                repeat = OxmlElement("w:tblHeader")   # repeat on each printed page
                repeat.set(qn("w:val"), "true")
                tr_pr.append(repeat)

    if snapshot_png:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        _run(p, labels["scan"], title_rtl, bold=True, size=9, gray=True)
        _para_dir(p, title_rtl)
        width = _picture_width(snapshot_png, usable)
        doc.add_picture(io.BytesIO(snapshot_png), width=width)
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    _run(p, labels["source"].format(file=source_name, p=page_no), title_rtl, size=8, gray=True)
    _para_dir(p, title_rtl)
    doc.core_properties.title = labels["table"].format(n=number)
    doc.core_properties.comments = "Created offline by HYPER-OCR"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def _base_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = _put(rpr, "rFonts", RPR_ORDER)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def _run(paragraph, text: str, rtl: bool, bold: bool = False, italic: bool = False, size: float | None = None,
         gray: bool = False):
    run = paragraph.add_run(text)
    run.bold = bold or None
    run.italic = italic or None
    if size:
        run.font.size = Pt(size)
    if gray:
        run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    rpr = run._r.get_or_add_rPr()
    if rtl:
        _put(rpr, "rtl", RPR_ORDER)
    if bold:
        _put(rpr, "bCs", RPR_ORDER)      # bold for Arabic (complex script) too
    if italic:
        _put(rpr, "iCs", RPR_ORDER)
    if size:
        _put(rpr, "szCs", RPR_ORDER, str(int(size * 2)))
    return run


def _para_dir(paragraph, rtl: bool) -> None:
    if rtl:
        _put(paragraph._p.get_or_add_pPr(), "bidi", PPR_ORDER)


def _picture_width(png: bytes, usable):
    from PIL import Image

    with Image.open(io.BytesIO(png)) as im:
        w_px = im.width
        dpi = (im.info.get("dpi") or (150, 150))[0] or 150
    natural = int(w_px / float(dpi) * 914400)   # EMU per inch
    return min(usable, natural)
