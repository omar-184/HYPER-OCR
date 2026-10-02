"""Whole conversions: the CPU engine on real scans, and the GPU engine's path with a
stand-in model that answers in Unlimited-OCR's documented output format."""

import zipfile

import pymupdf
import pytest
from docx import Document

from conftest import FIXTURES, needs_tesseract
from hyperocr import engines
from hyperocr.engines.base import Availability, Options
from hyperocr.engines.unlimited import UnlimitedOCREngine
from hyperocr.pipeline import convert

# What Unlimited-OCR would answer for the two English fixture pages (0-1000 grid).
FAKE_RAW = [
    "<|det|>title [93, 62, 680, 86]<|/det|>Effect of Early Mobilisation After Surgery\n"
    "<|det|>text [90, 94, 878, 141]<|/det|>Background: early mobilisation is recommended after abdominal surgery, "
    "but the evidence for its effect on length of stay is mixed. We reviewed 240 patients treated between 2021 and "
    "2023 at a single teaching hospital.\n"
    "<|det|>paragraph_title [91, 156, 179, 169]<|/det|>Methods\n"
    "<|det|>text [89, 180, 873, 212]<|/det|>Patients were grouped by the day of first walk. The primary outcome was "
    "length of hospital stay; secondary outcomes were readmission and pneumonia within 30 days.\n"
    "<|det|>image [264, 234, 758, 426]<|/det|>\n"
    "<|det|>figure_title [87, 454, 348, 466]<|/det|>Figure 1. Mean length of stay by group.\n"
    "<|det|>paragraph_title [87, 480, 170, 493]<|/det|>Results\n"
    "<|det|>text [87, 500, 722, 520]<|/det|>Patients who walked on day one went home sooner, with fewer complications.\n",
    "<|det|>table_title [83, 74, 420, 87]<|/det|>Table 1. Baseline characteristics\n"
    "<|det|>table [84, 98, 457, 217]<|/det|><table><tr><td>Group</td><td>Patients</td><td>Mean age</td>"
    "<td>Stay (days)</td></tr><tr><td>Day 1</td><td>96</td><td>54.2</td><td>4.1</td></tr><tr><td>Day 2</td>"
    "<td>88</td><td>57.9</td><td>5.3</td></tr><tr><td>Day 3+</td><td>56</td><td>61.4</td><td>7.8</td></tr></table>\n"
    "<|det|>text [85, 226, 821, 257]<|/det|>Readmission was lowest in the day-one group. The difference in pneumonia "
    "rates was not statistically significant \\(p = 0.08\\).\n"
    "<|det|>image [277, 262, 723, 498]<|/det|>\n"
    "<|det|>figure_title [87, 504, 395, 515]<|/det|>Figure 2. Chest radiograph of a typical patient.\n"
    "<|det|>paragraph_title [86, 530, 210, 543]<|/det|>Conclusion\n"
    "<|det|>text [86, 553, 700, 567]<|/det|>Walking on the first day after surgery is safe and shortens the hospital stay.\n",
]


class FakeUnlimited(UnlimitedOCREngine):
    """The real Unlimited-OCR engine class, with the model call replaced."""

    def __init__(self):
        super().__init__()
        self.calls = 0

    def availability(self):
        return Availability(True, detail="Test GPU")

    def prepare(self, progress):
        progress("loading-model")

    def _infer(self, image_path, out_dir):
        raw = FAKE_RAW[self.calls]
        self.calls += 1
        return raw


@pytest.fixture
def fake_gpu(monkeypatch):
    monkeypatch.setenv("HYPEROCR_GPU", "1")
    engines.get("tesseract")  # build the registry
    fake = FakeUnlimited()
    monkeypatch.setitem(engines._engines, "unlimited", fake)
    return fake


def _search(pdf, term):
    return sum(len(p.search_for(term)) for p in pymupdf.open(pdf))


@pytest.mark.gpu
def test_unlimited_path_end_to_end(tmp_path, fake_gpu):
    stages = []
    out = convert(FIXTURES / "scanned_english.pdf", tmp_path, Options(engine="auto"), lambda d: stages.append(d["stage"]))
    assert out.engine == "unlimited" and fake_gpu.calls == 2
    assert "loading-model" in stages and stages[-1] == "done"
    pdf = out.folder / out.pdf_file
    for term in ("Mobilisation", "pneumonia", "57.9", "radiograph"):
        assert _search(pdf, term), term
    # The title's text sits on the title's printed line (title box 62..86 of 1000 on a 842 pt page).
    hit = pymupdf.open(pdf)[0].search_for("Mobilisation")[0]
    assert 0.062 * 842 - 4 <= hit.y0 and hit.y1 <= 0.086 * 842 + 4
    assert out.images == ["Images/page-001_figure-01.png", "Images/page-002_figure-01.png"]
    assert [t["file"] for t in out.tables] == ["Tables/Table-01_page-002.docx"]
    doc = Document(str(out.folder / "Tables/Table-01_page-002.docx"))
    assert doc.tables[0].cell(2, 2).text == "57.9"
    assert "Table 1. Baseline characteristics" in [p.text for p in doc.paragraphs]
    md = out.markdown
    assert md.startswith("# Effect of Early Mobilisation After Surgery")
    assert "## Methods" in md and "![Figure 1. Mean length of stay by group.](Images/page-001_figure-01.png)" in md
    assert "| Day 2 | 88 | 57.9 | 5.3 |" in md
    assert "$p = 0.08$" in md
    toc = pymupdf.open(pdf).get_toc()
    assert toc[0][1] == "Effect of Early Mobilisation After Surgery"


@needs_tesseract
def test_tesseract_english(tmp_path):
    out = convert(FIXTURES / "scanned_english.pdf", tmp_path, Options(engine="tesseract", languages=["eng"]), lambda d: None)
    assert out.engine == "tesseract" and out.pages == 2
    pdf = out.folder / out.pdf_file
    for term in ("Mobilisation", "pneumonia", "Baseline", "57.9", "Conclusion"):
        assert _search(pdf, term), term
    assert len(out.images) == 2
    doc = Document(str(out.folder / out.tables[0]["file"]))
    cells = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert cells[0] == ["Group", "Patients", "Mean age", "Stay (days)"]
    assert cells[2] == ["Day 2", "88", "57.9", "5.3"]
    assert "## Methods" in out.markdown and "![" in out.markdown
    with zipfile.ZipFile(out.zip_path) as z:
        names = set(z.namelist())
    assert {"scanned_english/scanned_english_searchable.pdf", "scanned_english/scanned_english.md",
            "scanned_english/Images/page-001_figure-01.png", "scanned_english/Tables/Table-01_page-002.docx"} <= names


@needs_tesseract
def test_tesseract_arabic(tmp_path):
    langs = engines.get("tesseract").languages()
    if "ara" not in langs:
        pytest.skip("Arabic language data is not installed")
    out = convert(FIXTURES / "scanned_arabic.pdf", tmp_path, Options(engine="tesseract", languages=["eng", "ara"], ui_lang="ar"),
                  lambda d: None)
    pdf = out.folder / out.pdf_file
    for term in ("المريض", "الكرياتينين", "التوصيات", "Amlodipine", "13.5"):
        assert _search(pdf, term), term
    doc = Document(str(out.folder / out.tables[0]["file"]))
    rows = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert rows[0] == ["التحليل", "النتيجة", "المعدل الطبيعي"]
    assert rows[2][:2] == ["الكرياتينين", "1.1"]
    assert out.markdown.startswith("# تقرير المتابعة الطبية")
    # Tesseract's page reading skips this heading; the recovered line used to lose "نتا".
    assert "\n## نتائج التحاليل\n" in out.markdown
    assert _search(pdf, "نتائج")
    assert "الجدول 1 كملف Word" in out.markdown
    assert out.images == []


@needs_tesseract
def test_tesseract_arabic_small_print_setting(tmp_path):
    """At 400 dpi the page reading was confidently wrong on some Arabic words
    ("ااطبية", "سنئوات", "Jasall"); unsure words and cells are now re-read at a
    line height Tesseract reads well."""
    if "ara" not in engines.get("tesseract").languages():
        pytest.skip("Arabic language data is not installed")
    out = convert(FIXTURES / "scanned_arabic.pdf", tmp_path,
                  Options(engine="tesseract", languages=["eng", "ara"], dpi=400), lambda d: None)
    md = out.markdown
    for good in ("المتابعة الطبية", "خمس سنوات", "مراجعة نتائج التحاليل في العيادة", "## نتائج التحاليل"):
        assert good in md, good
    for bad in ("ااطبية", "سنئوات", "Jasall"):
        assert bad not in md, bad
    doc = Document(str(out.folder / out.tables[0]["file"]))
    rows = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert rows[0] == ["التحليل", "النتيجة", "المعدل الطبيعي"]
    assert rows[2] == ["الكرياتينين", "1.1", "1.2-0.6"]


@needs_tesseract
def test_rerunning_on_its_own_output_does_not_double_the_text(tmp_path):
    first = convert(FIXTURES / "scanned_english.pdf", tmp_path / "a", Options(engine="tesseract", languages=["eng"]), lambda d: None)
    second = convert(first.folder / first.pdf_file, tmp_path / "b", Options(engine="tesseract", languages=["eng"]), lambda d: None)
    assert _search(second.folder / second.pdf_file, "Mobilisation") == _search(first.folder / first.pdf_file, "Mobilisation")


def test_password_protected_pdf_is_refused(tmp_path):
    src = pymupdf.open(FIXTURES / "scanned_english.pdf")
    locked = tmp_path / "locked.pdf"
    src.save(locked, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="owner")
    with pytest.raises(engines.EngineError) as err:
        convert(locked, tmp_path, Options(), lambda d: None)
    assert err.value.key == "passwordProtected"
