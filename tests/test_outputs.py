"""Word tables, table grids and Markdown."""

from docx import Document
from docx.oxml.ns import qn

from hyperocr.document import CAPTION, FIGURE, FORMULA, HEADING, PAGE_NUMBER, TABLE, TEXT, TITLE, Block, PageResult
from hyperocr.outputs.markdown import to_markdown
from hyperocr.outputs.tablegrid import flat_html, parse_table
from hyperocr.outputs.tables import write_table_docx

SPANS = ('<table><tr><th>Group</th><th colspan="2">Patients</th></tr>'
         '<tr><td rowspan="2">Day 1</td><td>96</td><td>4.1</td></tr><tr><td>88</td><td>5.3</td></tr></table>')
ARABIC = ('<table dir="rtl"><tr><th>التحليل</th><th>النتيجة</th></tr>'
          '<tr><td>السكر الصائم</td><td>110</td></tr></table>')


def test_grid_resolves_spans():
    g = parse_table(SPANS)
    assert (g.rows, g.cols) == (3, 3)
    slots = g.at()
    assert slots[2][0].text == "Day 1" and slots[0][2].text == "Patients"


def test_flat_html_has_no_spans_and_repeats_row_headers():
    flat = flat_html(SPANS, "PIPE")
    assert "span" not in flat
    assert flat.count("Day 1") == 2


def test_docx_merges_and_header(tmp_path):
    path = tmp_path / "t.docx"
    write_table_docx(path, SPANS, 1, 2, "scan.pdf", caption="Table 1. Groups", snapshot_png=None)
    doc = Document(str(path))
    table = doc.tables[0]
    assert len(table.rows) == 3 and len(table.columns) == 3
    assert table.cell(0, 1)._tc is table.cell(0, 2)._tc     # colspan merged
    assert table.cell(1, 0)._tc is table.cell(2, 0)._tc     # rowspan merged
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert "Table 1. Groups" in [p.text for p in doc.paragraphs]


def test_docx_arabic_table_is_right_to_left(tmp_path):
    path = tmp_path / "ar.docx"
    write_table_docx(path, ARABIC, 1, 1, "scan.pdf", ui_lang="ar")
    doc = Document(str(path))
    assert doc.tables[0]._tbl.tblPr.find(qn("w:bidiVisual")) is not None
    assert doc.tables[0].cell(1, 0).text == "السكر الصائم"
    assert doc.paragraphs[0].text.startswith("جدول 1")


def test_markdown_structure():
    blocks = [
        Block(TITLE, (0, 0, 1, 1), text="A Study"),
        Block(HEADING, (0, 0, 1, 1), text="Methods", level=3),
        Block(TEXT, (0, 0, 1, 1), text="We used \\(x_1 + y\\) and snake_case."),
        Block(FIGURE, (0, 0, 1, 1), text="Figure 1. Chart", figure_file="Images/page-001_figure-01.png"),
        Block(CAPTION, (0, 0, 1, 1), text="Figure 1. Chart"),
        Block(TABLE, (0, 0, 1, 1), html="<table><tr><td>a|b</td><td>c</td></tr></table>", table_file="Tables/Table-01_page-001.docx"),
        Block(FORMULA, (0, 0, 1, 1), text="E = mc^2"),
        Block(PAGE_NUMBER, (0, 0, 1, 1), text="1"),
        Block(TEXT, (0, 0, 1, 1), text="يعاني المريض من ارتفاع ضغط الدم"),
    ]
    md = to_markdown([PageResult(0, 10, 10, 300, "test", blocks=blocks)], "A Study")
    assert md.startswith("# A Study\n")
    assert "\n## Methods\n" in md                            # levels renumbered from h2
    assert "$x_1 + y$" in md                                 # math untouched by escaping
    assert "![Figure 1. Chart](Images/page-001_figure-01.png)" in md
    assert "| a\\|b | c |" in md                              # pipe escaped inside the table
    assert "[Table 1 as a Word file](Tables/Table-01_page-001.docx)" in md
    assert "$$\nE = mc^2\n$$" in md
    assert "\n1\n" not in md                                  # page number left out
    assert "يعاني المريض من ارتفاع ضغط الدم" in md


def _assert_schema_order(path):
    """Every rPr, pPr and tblPr lists its children in the order the OOXML schema demands."""
    from hyperocr.outputs.tables import PPR_ORDER, RPR_ORDER, TBLPR_ORDER

    doc = Document(str(path))
    body = doc.element.body
    checked = 0
    for tag, order in (("w:rPr", RPR_ORDER), ("w:pPr", PPR_ORDER), ("w:tblPr", TBLPR_ORDER)):
        for el in body.iter(qn(tag)):
            names = [c.tag.rsplit("}", 1)[-1] for c in el]
            ranks = [order.index(n) for n in names if n in order]
            assert ranks == sorted(ranks), (tag, names)
            checked += 1
    assert checked > 5


def test_docx_elements_follow_schema_order(tmp_path):
    for html, lang in ((SPANS, "en"), (ARABIC, "ar")):
        path = tmp_path / ("t-%s.docx" % lang)
        import io
        from PIL import Image
        buf = io.BytesIO()
        Image.new("RGB", (300, 120), "white").save(buf, "PNG")
        write_table_docx(path, html, 1, 1, "scan.pdf", caption="Table 1", snapshot_png=buf.getvalue(), ui_lang=lang)
        _assert_schema_order(path)
