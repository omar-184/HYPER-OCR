"""Pages made on a computer: their own text is used, no OCR. And the files a conversion makes
are the ones chosen."""

import zipfile

import numpy as np
import pymupdf
import pytest
from docx import Document

from conftest import FIXTURES, needs_tesseract
from hyperocr import engines
from hyperocr.engines import pdftext
from hyperocr.engines.base import Options
from hyperocr.pipeline import convert

ROWS = [["Group", "Patients", "Mean age", "Stay (days)"], ["Day 1", "96", "54.2", "4.1"],
        ["Day 2", "88", "57.9", "5.3"], ["Day 3+", "56", "61.4", "7.8"]]
BODY = ("Background: early mobilisation is recommended after abdominal surgery, but the evidence for its "
        "effect on length of stay is mixed. We reviewed 240 patients treated between 2021 and 2023 at a "
        "single teaching hospital, and grouped them by the day of their first walk.")


def made_on_a_computer(path, rotation=0):
    """A one-page report as Word would export it: real text, a bordered table, a chart."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 90), "Effect of Early Mobilisation After Surgery", fontsize=20, fontname="Helvetica-Bold")
    page.insert_textbox(pymupdf.Rect(72, 110, 523, 200), BODY, fontsize=11, fontname="Helvetica")
    page.insert_text((72, 230), "Results", fontsize=15, fontname="Helvetica-Bold")
    x, y, w, h = 72, 250, 110, 26
    for r, row in enumerate(ROWS):
        for c, text in enumerate(row):
            cell = pymupdf.Rect(x + c * w, y + r * h, x + (c + 1) * w, y + (r + 1) * h)
            page.draw_rect(cell, color=(0, 0, 0), width=0.8)
            page.insert_text((cell.x0 + 6, cell.y1 - 9), text, fontsize=10, fontname="Helvetica")
    chart = np.full((300, 500, 3), 255, np.uint8)               # a bar chart, as a picture
    for i, height in enumerate((180, 120, 220, 90)):
        chart[280 - height:280, 40 + i * 110:110 + i * 110] = (40, 90, 200)
    chart[280:284, 20:480] = 0
    pix = pymupdf.Pixmap(pymupdf.csRGB, 500, 300, chart.tobytes(), False)
    page.insert_image(pymupdf.Rect(72, 400, 372, 580), pixmap=pix)
    page.insert_text((72, 600), "Figure 1. Mean length of stay by group.", fontsize=10, fontname="Helvetica-Oblique")
    page.set_rotation(rotation)
    doc.save(path)
    return path


@pytest.fixture
def no_ocr(monkeypatch):
    """Fails any test that reaches Tesseract."""
    import pytesseract.pytesseract as ptp

    def refuse(*args, **kwargs):
        raise AssertionError("Tesseract was run")

    monkeypatch.setattr(ptp, "run_tesseract", refuse)


def test_a_page_made_on_a_computer_is_taken_as_it_is(tmp_path, no_ocr):
    out = convert(made_on_a_computer(tmp_path / "report.pdf"), tmp_path / "out",
                  Options(engine="tesseract", languages=["eng"]), lambda d: None)
    assert out.engine == "pdf-text"
    assert {"key": "ownText", "pages": [1]} in out.warnings
    assert not any(w["key"] == "keptText" for w in out.warnings)       # said once
    md = out.markdown
    assert md.startswith("# Effect of Early Mobilisation After Surgery")
    assert "## Results" in md
    assert BODY in md                                                    # exactly, every word
    assert "| Day 2 | 88 | 57.9 | 5.3 |" in md
    doc = Document(str(out.folder / out.tables[0]["file"]))
    assert [[c.text for c in row.cells] for row in doc.tables[0].rows] == ROWS
    assert len(out.images) == 1 and "*Figure 1. Mean length of stay by group.*" in md
    pdf = pymupdf.open(out.folder / out.pdf_file)
    assert len(pdf[0].search_for("Background")) == 1                     # its own text, not a second copy


def test_a_turned_page_made_on_a_computer_lines_up(tmp_path, no_ocr):
    """/Rotate turns the page as shown; the words must land where they are shown."""
    path = made_on_a_computer(tmp_path / "turned.pdf", rotation=90)
    page = pymupdf.open(path)[0]
    dpi = 100
    words = pdftext.page_words(page, dpi)
    pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY)
    image = np.frombuffer(pix.samples, np.uint8).reshape(pix.h, pix.w)
    title = next(w for w in words if w.text == "Mobilisation")
    x0, y0, x1, y1 = (int(v) for v in title.box)
    assert (image[y0:y1, x0:x1] < 128).mean() > 0.05                    # ink under the word
    assert y1 - y0 > x1 - x0                                             # running down the turned page


def test_scans_and_pages_in_arabic_are_still_read_by_ocr(tmp_path):
    scan = pymupdf.open(FIXTURES / "scanned_english.pdf")[0]
    assert pdftext.page_words(scan, 300) is None
    assert not pdftext._usable("تقرير المتابعة الطبية للمريض " * 5)         # right-to-left: OCR
    assert not pdftext._usable("Th� p�tient ���")  # undecodable: OCR
    assert pdftext._usable(BODY)


@needs_tesseract
def test_the_switch_turns_it_off(tmp_path):
    out = convert(made_on_a_computer(tmp_path / "report.pdf"), tmp_path / "out",
                  Options(engine="tesseract", languages=["eng"], own_text=False), lambda d: None)
    assert out.engine == "tesseract"
    assert {"key": "keptText", "pages": [1]} in out.warnings and not any(w["key"] == "ownText" for w in out.warnings)


def test_only_the_chosen_files_are_made(tmp_path, no_ocr):
    path = made_on_a_computer(tmp_path / "report.pdf")
    out = convert(path, tmp_path / "md", Options(engine="tesseract", outputs=("markdown",)), lambda d: None)
    assert out.pdf_file == "" and out.md_file == "report.md" and out.images == [] and out.tables == []
    assert "| Day 2 | 88 | 57.9 | 5.3 |" in out.markdown                 # the table is still in the Markdown
    with zipfile.ZipFile(out.zip_path) as z:
        assert z.namelist() == ["report/report.md"]

    out = convert(path, tmp_path / "pdf", Options(engine="tesseract", outputs=("pdf", "tables")), lambda d: None)
    assert out.md_file == "" and out.markdown == "" and out.images == []
    with zipfile.ZipFile(out.zip_path) as z:
        names = set(z.namelist())
    assert names == {"report/report_searchable.pdf", "report/Tables/", "report/Tables/Table-01_page-001.docx"}


def test_the_server_takes_the_choices(tmp_path):
    from hyperocr.server import _options

    info = {"languages": ["eng"]}
    assert _options({"outputs": ["tables", "pdf", "nonsense"]}, info).outputs == ("pdf", "tables")
    assert _options({"outputs": []}, info).outputs == ("pdf", "markdown", "images", "tables")
    assert _options({}, info).own_text is True and _options({"own_text": False}, info).own_text is False
    engines.get("tesseract")    # the registry still builds
